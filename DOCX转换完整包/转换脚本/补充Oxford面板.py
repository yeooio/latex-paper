#!/usr/bin/env python3
"""
补充插入Oxford面板图片
（在现有DOCX基础上继续插入）
"""
from pathlib import Path
from docx import Document
from docx.shared import Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_DIR = Path(__file__).parent
CONVERSION_DIR = PROJECT_DIR / "output" / "final_docx_conversion"
IMAGES_PROCESSED = CONVERSION_DIR / "images_processed"

def insert_oxford_panel():
    """在现有DOCX中插入Oxford面板"""

    # 找到最新的FINAL文件
    final_files = sorted(CONVERSION_DIR.glob("FINAL_thesis_complete_*.docx"))
    if not final_files:
        print("✗ 找不到FINAL文件")
        return

    latest_docx = final_files[-1]
    print(f"正在处理: {latest_docx.name}")

    # 检查Oxford面板是否存在
    oxford_panel = IMAGES_PROCESSED / "figure_4_2_oxford_panel.png"
    if not oxford_panel.exists():
        print(f"✗ Oxford面板还未生成: {oxford_panel.name}")
        return

    print(f"✓ 找到Oxford面板: {oxford_panel.name}")

    # 打开DOCX
    doc = Document(str(latest_docx))

    # 查找插入位置（关键词匹配）
    keywords = [
        'SOH estimation results',
        'Oxford Cell',
        'estimation errors of different',
    ]

    inserted = False
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        for keyword in keywords:
            if keyword.lower() in text.lower():
                try:
                    # 插入图片
                    new_para = para.insert_paragraph_before()
                    run = new_para.add_run()
                    run.add_picture(str(oxford_panel), width=Mm(165))
                    new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    print(f"✓ 已插入Oxford面板")
                    print(f"  位置: 第{i}段")
                    print(f"  匹配: {keyword}")
                    inserted = True
                    break
                except Exception as e:
                    print(f"✗ 插入失败: {e}")

        if inserted:
            break

    if not inserted:
        print("⚠ 未找到合适的插入位置")
        return

    # 保存更新的文件
    import datetime
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    output_docx = CONVERSION_DIR / f"FINAL_thesis_complete_ALL_{timestamp}.docx"
    doc.save(str(output_docx))

    size_mb = output_docx.stat().st_size / 1024 / 1024
    print(f"\n✓ 已保存更新文件:")
    print(f"  文件: {output_docx.name}")
    print(f"  大小: {size_mb:.2f} MB")
    print(f"\n现在包含所有11个图片！")

if __name__ == "__main__":
    print("="*70)
    print("补充插入Oxford面板")
    print("="*70)
    insert_oxford_panel()
