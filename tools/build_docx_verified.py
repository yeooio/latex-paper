"""Deterministic LaTeX -> Markdown -> DOCX with editable math and PDF table crops."""
from pathlib import Path
import re, json, subprocess, hashlib, sys, copy
import pdfplumber
from PIL import Image
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'docx_verified'
OUT.mkdir(parents=True, exist_ok=True)
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
PANDOC = ROOT / 'tmp/docx-deps/pypandoc/files/pandoc.exe'
PDF = ROOT / 'main.pdf'
manifest = {'source_pdf': str(PDF), 'pdf_sha256': hashlib.sha256(PDF.read_bytes()).hexdigest(), 'tables': [], 'figures': [], 'equations': []}

def group(text, start):
    assert text[start] == '{'
    depth = 1
    for i in range(start + 1, len(text)):
        if text[i] in '{}' and (i == 0 or text[i-1] != '\\'):
            depth += 1 if text[i] == '{' else -1
            if depth == 0: return text[start+1:i], i+1
    raise ValueError('Unbalanced braces')

def arg(text, cmd):
    m = re.search(r'\\' + cmd + r'(?:\[[^\]]*\])?\s*\{', text)
    return group(text, m.end()-1)[0] if m else ''

def bookmark_name(label): return 'b_' + re.sub('[^A-Za-z0-9_]', '_', label)

labels = {}
aux = (ROOT / 'main.aux').read_text(encoding='utf-8')
for line in aux.splitlines():
    m = re.match(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}', line)
    if m and '@cref' not in m[1]: labels[m[1]] = m[2]

def clean_comments(text): return re.sub(r'(?<!\\)%[^\n]*', '', text)

def expand(path):
    text = clean_comments(path.read_text(encoding='utf-8'))
    def repl(m):
        rel = m[1]
        p = ROOT / (rel if rel.endswith('.tex') else rel+'.tex')
        if rel == 'preamble': return ''
        if rel.startswith('tables/'):
            s = clean_comments(p.read_text(encoding='utf-8'))
            label = arg(s, 'label'); number = int(labels[label])
            manifest['tables'].append({'number': number, 'label': label, 'caption': arg(s,'caption'), 'source': str(p)})
            return '\n\nDOCXTABLE%02d\n\n' % number
        if rel.startswith('figures/'):
            s = clean_comments(p.read_text(encoding='utf-8'))
            label = arg(s, 'label'); number = int(labels[label])
            images = re.findall(r'\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}', s)
            manifest['figures'].append({'number':number,'label':label,'caption':arg(s,'caption'), 'images':images,'source':str(p)})
            return '\n\nDOCXFIG%02d\n\n' % number
        return '\n' + expand(p) + '\n'
    return re.sub(r'\\input\{([^}]+)\}', repl, text)

text = expand(ROOT / 'main.tex')
text = text.split('\\begin{document}',1)[1].rsplit('\\end{document}',1)[0]
# Replace bibliography with explicit anchors, keeping every entry and DOI.
bib = re.search(r'\\begin\{thebibliography\}\{[^}]+\}(.*?)\\end\{thebibliography\}', text, re.S)
assert bib
entries = re.findall(r'\\bibitem\{([^}]+)\}\s*(.*?)(?=\\bibitem|\Z)',bib[1],re.S)
bib_text = '\n\\section*{参考文献}\n'
for label, body in entries:
    body = body.strip()
    number = re.search(r'\d+',label)[0]
    bib_text += '\nDOCXBIB%s\n\n[%s] %s\n\n' % (number,number,body)
text = text[:bib.start()] + bib_text + text[bib.end():]

def citation(m):
    ids = m[1].split(',')
    nums = sorted(int(re.search(r'\d+',x)[0]) for x in ids)
    parts=[]; j=0
    while j<len(nums):
        k=j
        while k+1<len(nums) and nums[k+1]==nums[k]+1:k+=1
        if k-j>=2:
            parts += ['\\href{#b_ref%d}{%d}--\\href{#b_ref%d}{%d}' % (nums[j],nums[j],nums[k],nums[k])]
        else:
            parts += ['\\href{#b_ref%d}{%d}' % (v,v) for v in nums[j:k+1]]
        j=k+1
    return '[' + ','.join(parts) + ']'
text = re.sub(r'\\cite\{([^}]+)\}',citation,text)

def crossref(m):
    cmd, label = m[1],m[2]
    if label not in labels: raise ValueError('Unknown reference '+label)
    value = labels[label]
    if cmd == 'eqref': value='('+value+')'
    elif cmd in ('cref','Cref'):
        prefix = 'Fig. ' if label.startswith('fig:') else 'Table ' if label.startswith('tab:') else 'Eq. ' if label.startswith('eq:') else ''
        value = prefix + ('('+value+')' if label.startswith('eq:') else value)
    return '\\href{#%s}{%s}' % (bookmark_name(label),value)
text = re.sub(r'\\(eqref|cref|Cref|ref)\{([^}]+)\}',crossref,text)

eqn = 0
def equation(m):
    global eqn
    star,body=m[1],m[2]
    labs=re.findall(r'\\label\{([^}]+)\}',body)
    body=re.sub(r'\\label\{[^}]+\}','',body).strip()
    if star: return '\\['+body+'\\]'
    eqn+=1
    for lab in labs: assert labels[lab]==str(eqn),(lab,labels[lab],eqn)
    manifest['equations'].append({'number':eqn,'labels':labs,'latex':body})
    return '\\['+body+'\\]\n\nDOCXEQ%02d\n\n' % eqn
text = re.sub(r'\\begin\{equation(\*?)\}(.*?)\\end\{equation\1\}',equation,text,flags=re.S)
text = re.sub(r'\\label\{([^}]+)\}', lambda m:'\n\nDOCXANCHOR'+bookmark_name(m[1])+'\n\n',text)
text = re.sub(r'\\fontsize\{[^}]+\}\{[^}]+\}\\selectfont','',text)
text = re.sub(r'\\setlength\{[^}]+\}\{[^}]+\}','',text)
text = re.sub(r'\\(?:FloatBarrier|begingroup|endgroup|sloppy|allowbreak)\b','',text)
(OUT/'conversion.tex').write_text(text,encoding='utf-8')
def pandoc(*args):
    r=subprocess.run([str(PANDOC),*map(str,args)],cwd=ROOT,capture_output=True,encoding='utf-8')
    if r.stderr: print(r.stderr)
    r.check_returncode()
pandoc(OUT/'conversion.tex','-f','latex','-t','markdown','-o',OUT/'paper.md','--wrap=none')
pandoc(OUT/'paper.md','-f','markdown','-t','docx','-o',OUT/'base.docx','--number-sections')

# Obtain exact table boundaries from the caption plus booktabs rules.
with pdfplumber.open(PDF) as pdf:
    tables={}
    for page_no,page in enumerate(pdf.pages,1):
        for word in page.extract_words():
            m=re.fullmatch(r'Table\s*(\d+)',word['text'])
            if not m or word['x0']>100: continue
            n=int(m[1]); top=word['top']
            rules=sorted(set(round(e['top'],2) for e in page.edges if e['x1']-e['x0']>450 and e['height']<2 and e['top']>top))
            assert len(rules)>=3,(n,rules)
            bottom=rules[2]+4
            if n in (11,12): bottom += 14 # test-cell explanatory footnote
            box=(54,top-2,541,bottom)
            im=page.crop(box).to_image(resolution=320).original
            dest=ASSETS/f'table_{n:02d}.png'; im.save(dest)
            tables[n]={'image':str(dest),'page':page_no,'bbox':box}
    assert len(tables)==len(manifest['tables'])==18,(len(tables),len(manifest['tables']))
    for t in manifest['tables']: t.update(tables[t['number']])

doc=Document(OUT/'base.docx')
sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.left_margin=sec.right_margin=Mm(20.5)
sec.top_margin=sec.bottom_margin=Mm(22)
sec.footer_distance=Mm(10)

def fonts(style, name='宋体', size=11):
    style.font.name=name;style.font.size=Pt(size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)
    style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
for s in doc.styles:
    if s.type==1:
        fonts(s)
        s.paragraph_format.space_after=Pt(0)
        s.paragraph_format.line_spacing=1.05
for name in ['Normal','Body Text','First Paragraph']:
    if name in doc.styles:
        s=doc.styles[name];s.paragraph_format.first_line_indent=Pt(22)
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for i,size in [(1,14),(2,13),(3,12)]:
    s=doc.styles['Heading '+str(i)];fonts(s,'黑体',size);s.font.bold=True
    s.paragraph_format.space_before=Pt(6);s.paragraph_format.space_after=Pt(4)
    s.paragraph_format.keep_with_next=True;s.paragraph_format.first_line_indent=Pt(0)

bid=1000
def bookmark(p,name):
    global bid
    bid+=1
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bid));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bid))
    p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end)

def delete(p):p._p.getparent().remove(p._p)
def set_image(p,path,width=169):
    p.clear();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(4)
    p.add_run().add_picture(str(path),width=Mm(width))
    p.paragraph_format.keep_together=True

def plain_caption(tex):
    r=subprocess.run([str(PANDOC),'-f','latex','-t','plain'],input=tex,capture_output=True,encoding='utf-8')
    r.check_returncode();return r.stdout.strip().replace('\n',' ')

for p in list(doc.paragraphs):
    s=p.text.strip()
    if s.startswith('DOCXTABLE'):
        n=int(s[9:]);t=next(x for x in manifest['tables'] if x['number']==n)
        set_image(p,t['image']);bookmark(p,bookmark_name(t['label']))
    elif s.startswith('DOCXFIG'):
        n=int(s[7:]);f=next(x for x in manifest['figures'] if x['number']==n)
        if n in (8,9):
            path=ROOT/'DOCX转换完整包/处理后的图片'/('figure_4_2_oxford_panel.png' if n==8 else 'figure_4_3_calce_mit_panel.png')
        elif n==10:
            with pdfplumber.open(ROOT/f['images'][0]) as figpdf:
                path=ASSETS/'figure_10.png';figpdf.pages[0].to_image(resolution=300).save(path)
        else:path=ROOT/f['images'][0]
        f['image']=str(path)
        caption=p.insert_paragraph_before();p._p.addnext(caption._p)
        set_image(p,path,169 if n not in (5,10) else 155)
        bookmark(p,bookmark_name(f['label']));p.paragraph_format.keep_with_next=True
        caption.text='Fig. %d. %s'%(n,plain_caption(f['caption']))
        caption.alignment=WD_ALIGN_PARAGRAPH.CENTER;caption.paragraph_format.first_line_indent=Pt(0)
        for r in caption.runs:r.font.size=Pt(9)
    elif s.startswith('DOCXEQ'):
        n=int(s[6:]);e=manifest['equations'][n-1]
        prev=p._p.getprevious()
        assert prev is not None and any(x.tag == qn('m:oMath') for x in prev.iter()),n
        from docx.text.paragraph import Paragraph
        ep=Paragraph(prev,p._parent)
        # Flatten display wrapper so equation and number share one Word paragraph.
        for mp in list(prev):
            if mp.tag != qn('m:oMathPara'): continue
            for math in list(mp):
                if math.tag == qn('m:oMath'): prev.insert(prev.index(mp),math)
            prev.remove(mp)
        ep.paragraph_format.first_line_indent=Pt(0);ep.alignment=WD_ALIGN_PARAGRAPH.LEFT
        ep.paragraph_format.tab_stops.add_tab_stop(Mm(84.5),WD_TAB_ALIGNMENT.CENTER)
        ep.paragraph_format.tab_stops.add_tab_stop(Mm(169),WD_TAB_ALIGNMENT.RIGHT)
        run=OxmlElement('w:r');run.append(OxmlElement('w:tab'));prev.insert(1 if prev.pPr is not None else 0,run)
        ep.add_run('\t(%d)'%n)
        ep.paragraph_format.space_before=Pt(5);ep.paragraph_format.space_after=Pt(5)
        ep.paragraph_format.keep_together=True
        for lab in e['labels']:bookmark(ep,bookmark_name(lab))
        bookmark(ep,'equation_%d'%n);delete(p)
    elif s.startswith('DOCXBIB'):
        from docx.text.paragraph import Paragraph
        n=int(s[7:]);np=Paragraph(p._p.getnext(),p._parent)
        bookmark(np,'b_ref%d'%n)
        np.paragraph_format.first_line_indent=Pt(-18);np.paragraph_format.left_indent=Pt(18)
        np.paragraph_format.space_after=Pt(2);delete(p)
    elif s.startswith('DOCXANCHOR'):
        from docx.text.paragraph import Paragraph
        bookmark(Paragraph(p._p.getprevious(),p._parent),s[10:]);delete(p)

# Footer page field and equation font.
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for math in doc._element.xpath('.//m:oMath'):
    for r in math.iter(qn('m:r')):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        ft=OxmlElement('w:rFonts');ft.set(qn('w:ascii'),'Cambria Math');ft.set(qn('w:hAnsi'),'Cambria Math');pr.append(ft)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');pr.append(sz)
settings=doc.settings.element
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)

destination=OUT/'MS-AgentNet_可编辑论文.docx'
doc.save(destination)
manifest.update({'destination':str(destination),'references':len(entries),'numbered_equations':eqn,'images':len(doc.inline_shapes),'status':'awaiting_render_verification'})
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k not in ('tables','figures','equations')},ensure_ascii=False,indent=2))
